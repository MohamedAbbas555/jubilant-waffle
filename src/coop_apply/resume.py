"""Read a resume (PDF, DOCX, TXT/MD) and turn it into a structured ``ResumeProfile``."""

from __future__ import annotations

import base64
from pathlib import Path

from .llm import LLM
from .models import ResumeProfile
from .utils import file_hash, normalize_whitespace

PARSE_SYSTEM = """You extract a structured profile from a university co-op student's resume.
Use only information that is present in the resume. Leave a field as an empty string or
empty list when the resume doesn't state it; never guess contact details. Keep experience
highlights close to the original wording (lightly condensed), since they will later be
quoted in cover letters. For target_roles, infer the 3-6 kinds of co-op roles the resume
is most competitive for (e.g. "backend software developer", "data analyst",
"embedded firmware developer")."""


def extract_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif suffix in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported resume format: {suffix} (use .pdf, .docx, .txt or .md)")
    return normalize_whitespace(text)


def parse_resume(path: str | Path, llm: LLM, effort: str = "medium") -> ResumeProfile:
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        # Send the PDF itself so Claude sees layout (columns, headings) that text
        # extraction often scrambles.
        data = base64.standard_b64encode(path.read_bytes()).decode()
        content = [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data}},
            {"type": "text", "text": "Extract the profile from this resume."},
        ]
    else:
        content = f"<resume>\n{extract_text(path)}\n</resume>\n\nExtract the profile from this resume."
    return llm.structured(system=PARSE_SYSTEM, content=content, schema=ResumeProfile, effort=effort)


class ProfileCache:
    """Caches the parsed profile next to the database, keyed by the resume's hash."""

    def __init__(self, data_dir: Path):
        self.dir = data_dir / "profiles"

    def path_for(self, resume_hash: str) -> Path:
        return self.dir / f"{resume_hash[:16]}.json"

    def load(self, resume_path: str | Path) -> tuple[ResumeProfile | None, str]:
        h = file_hash(resume_path)
        p = self.path_for(h)
        if p.exists():
            return ResumeProfile.model_validate_json(p.read_text()), h
        return None, h

    def save(self, resume_hash: str, profile: ResumeProfile) -> Path:
        self.dir.mkdir(parents=True, exist_ok=True)
        p = self.path_for(resume_hash)
        p.write_text(profile.model_dump_json(indent=2))
        return p


def load_or_parse(resume_path: str | Path, data_dir: Path, llm: LLM, effort: str = "medium", force: bool = False):
    if not Path(resume_path).exists():
        raise FileNotFoundError(f"Resume not found: {resume_path} (set profile.resume_path in config.toml)")
    cache = ProfileCache(data_dir)
    profile, h = cache.load(resume_path)
    if profile is None or force:
        profile = parse_resume(resume_path, llm, effort=effort)
        cache.save(h, profile)
    return profile, h
